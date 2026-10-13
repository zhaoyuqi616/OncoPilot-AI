from __future__ import annotations

import hashlib
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from oncopilot_ai.config import settings
from oncopilot_ai.security import detect_prompt_injection


def _sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as fh:
        for block in iter(lambda: fh.read(1024 * 1024), b""):
            h.update(block)
    return h.hexdigest()


def _base_metadata(path: Path, workspace_id: str, defaults: dict[str, Any] | None) -> dict[str, Any]:
    defaults = defaults or {}
    return {
        "workspace_id": workspace_id,
        "source_filename": path.name,
        "source_sha256": _sha256(path),
        "sensitivity": str(defaults.get("sensitivity", "internal")).lower(),
        "allowed_roles": list(defaults.get("allowed_roles", [])),
        "allowed_users": list(defaults.get("allowed_users", [])),
        "ingested_at": datetime.now(timezone.utc).isoformat(),
    }


def _load_txt(path: Path, workspace_id: str, defaults: dict[str, Any] | None) -> list[dict[str, Any]]:
    text = path.read_text(encoding="utf-8", errors="replace")
    meta = _base_metadata(path, workspace_id, defaults)
    return [{
        "doc_id": meta["source_sha256"][:16],
        "title": path.stem.replace("_", " ").title(),
        "text": text,
        "section": "document",
        "parser": "text",
        "ocr_used": False,
        "prompt_injection_suspected": detect_prompt_injection(text),
        **meta,
    }]


def _load_docx(path: Path, workspace_id: str, defaults: dict[str, Any] | None) -> list[dict[str, Any]]:
    from docx import Document

    doc = Document(path)
    paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                paragraphs.append(" | ".join(cells))
    text = "\n".join(paragraphs)
    meta = _base_metadata(path, workspace_id, defaults)
    return [{
        "doc_id": meta["source_sha256"][:16],
        "title": path.stem.replace("_", " ").title(),
        "text": text,
        "section": "document",
        "parser": "python-docx",
        "ocr_used": False,
        "prompt_injection_suspected": detect_prompt_injection(text),
        **meta,
    }]


def _ocr_page(page) -> str:
    if not settings.ocr_enabled:
        return ""
    try:
        import pytesseract
        from PIL import Image
    except ImportError as exc:  # pragma: no cover
        raise RuntimeError("OCR requested but Pillow/pytesseract are not installed") from exc
    pix = page.get_pixmap(matrix=__import__("fitz").Matrix(2, 2), alpha=False)
    image = Image.frombytes("RGB", [pix.width, pix.height], pix.samples)
    return pytesseract.image_to_string(image)


def _load_pdf(path: Path, workspace_id: str, defaults: dict[str, Any] | None) -> list[dict[str, Any]]:
    import fitz

    pdf = fitz.open(path)
    meta = _base_metadata(path, workspace_id, defaults)
    records: list[dict[str, Any]] = []
    for page_number, page in enumerate(pdf, start=1):
        native_text = page.get_text("text").strip()
        ocr_used = len(native_text) < settings.ocr_min_native_chars_per_page
        text = _ocr_page(page).strip() if ocr_used else native_text
        if not text:
            continue
        records.append({
            "doc_id": meta["source_sha256"][:16],
            "title": path.stem.replace("_", " ").title(),
            "text": text,
            "page_number": page_number,
            "section": f"page_{page_number}",
            "parser": "pymupdf+ocr" if ocr_used else "pymupdf",
            "ocr_used": ocr_used,
            "prompt_injection_suspected": detect_prompt_injection(text),
            **meta,
        })
    return records


def load_document(path: str | Path, workspace_id: str, defaults: dict[str, Any] | None = None) -> list[dict[str, Any]]:
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix not in settings.allowed_extension_set:
        raise ValueError(f"Unsupported file type: {suffix}")
    if path.stat().st_size > settings.max_upload_mb * 1024 * 1024:
        raise ValueError(f"File exceeds {settings.max_upload_mb} MB limit")
    if suffix == ".txt":
        return _load_txt(path, workspace_id, defaults)
    if suffix == ".docx":
        return _load_docx(path, workspace_id, defaults)
    if suffix == ".pdf":
        return _load_pdf(path, workspace_id, defaults)
    raise ValueError(f"Unsupported file type: {suffix}")


def load_documents_from_workspace(workspace_id: str, defaults: dict[str, Any] | None = None) -> tuple[list[dict[str, Any]], list[dict[str, str]]]:
    incoming = Path(settings.workspace_data_root) / workspace_id / "incoming"
    incoming.mkdir(parents=True, exist_ok=True)
    documents: list[dict[str, Any]] = []
    errors: list[dict[str, str]] = []
    seen_hashes: set[str] = set()
    for path in sorted(p for p in incoming.iterdir() if p.is_file() and p.suffix.lower() in settings.allowed_extension_set):
        try:
            records = load_document(path, workspace_id, defaults)
            if records and records[0]["source_sha256"] in seen_hashes:
                continue
            if records:
                seen_hashes.add(records[0]["source_sha256"])
            documents.extend(records)
        except Exception as exc:
            errors.append({"file": path.name, "error": str(exc)})
    return documents, errors
